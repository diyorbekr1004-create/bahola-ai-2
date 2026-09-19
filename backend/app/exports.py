"""Eksportlar: Excel (to'liq), HEMIS-format (xlsx/csv), DOCX (silabus, hisobot, talaba fikr-mulohazasi)."""
from __future__ import annotations

import csv
import io
import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import config
from .analytics import GRADE_LABELS, percent, score_to_grade5

DEFAULT_HEMIS_TEMPLATE: Dict[str, Any] = {
    "sheet_name": "Baholar",
    "date_format": "%d.%m.%Y",
    "columns": [
        {"field": "row_number", "header": "№"},
        {"field": "student_id", "header": "Talaba ID (HEMIS)"},
        {"field": "student_name", "header": "F.I.Sh."},
        {"field": "group_name", "header": "Guruh"},
        {"field": "course", "header": "Fan"},
        {"field": "topic", "header": "Mavzu / Topshiriq"},
        {"field": "control_type", "header": "Nazorat turi"},
        {"field": "score", "header": "Ball (100)"},
        {"field": "grade_5", "header": "Baho (5)"},
        {"field": "date", "header": "Sana"},
        {"field": "teacher", "header": "O'qituvchi"},
        {"field": "status", "header": "Holat"},
    ],
    "status_labels": {"pending": "Tasdiqlanmagan", "confirmed": "Tasdiqlangan", "edited": "Tuzatilgan"},
}

HEADER_FILL = PatternFill("solid", fgColor="2A78D6")
HEADER_FONT = Font(bold=True, color="FFFFFF")
FAIL_FILL = PatternFill("solid", fgColor="FBE3E3")
GOOD_FILL = PatternFill("solid", fgColor="E3F4E3")
THIN = Side(style="thin", color="C3C2B7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


# ---------------------------------------------------------------------------
# Umumiy
# ---------------------------------------------------------------------------
def export_dir() -> Path:
    d = config.settings.export_dir
    d.mkdir(parents=True, exist_ok=True)
    return d


def safe_name(value: Optional[str], default: str = "export") -> str:
    value = (value or default).strip()
    value = re.sub(r"[^\w\-]+", "_", value, flags=re.UNICODE).strip("_")
    return value[:60] or default


def timestamp() -> str:
    return datetime.utcnow().strftime("%Y%m%d_%H%M%S")


def make_path(prefix: str, label: Optional[str], ext: str) -> Path:
    return export_dir() / f"{prefix}_{safe_name(label)}_{timestamp()}.{ext}"


def load_hemis_template() -> Dict[str, Any]:
    path = config.settings.hemis_template_path
    try:
        with open(path, "r", encoding="utf-8") as fh:
            data = json.load(fh)
        if isinstance(data, dict) and isinstance(data.get("columns"), list) and data["columns"]:
            merged = dict(DEFAULT_HEMIS_TEMPLATE)
            merged.update({k: v for k, v in data.items() if not k.startswith("_")})
            return merged
    except (OSError, ValueError):
        pass
    return dict(DEFAULT_HEMIS_TEMPLATE)


def _style_header(ws, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.freeze_panes = "A2"


def _autosize(ws, min_width: int = 8, max_width: int = 45) -> None:
    widths: Dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            widths[cell.column] = max(widths.get(cell.column, 0), len(str(cell.value)))
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(min_width, min(max_width, w + 2))


def _fmt_date(value: Any, fmt: str) -> str:
    if not value:
        return ""
    if isinstance(value, datetime):
        return value.strftime(fmt)
    try:
        return datetime.fromisoformat(str(value)).strftime(fmt)
    except ValueError:
        return str(value)


# ---------------------------------------------------------------------------
# HEMIS eksport
# ---------------------------------------------------------------------------
def hemis_rows(rows: List[Dict[str, Any]], control_type: str = "JN", template: Optional[Dict[str, Any]] = None) -> List[Dict[str, Any]]:
    tpl = template or load_hemis_template()
    labels = tpl.get("status_labels", {})
    out: List[Dict[str, Any]] = []
    for i, r in enumerate(rows, start=1):
        score_pct = percent(r.get("total_score"), r.get("max_score") or 100)
        status = r.get("status") or ("confirmed" if r.get("confirmed") else "pending")
        out.append({
            "row_number": i,
            "student_id": r.get("student_id") or "",
            "student_name": r.get("student_name") or "",
            "group_name": r.get("group_name") or "",
            "course": r.get("course") or "",
            "topic": r.get("topic") or "",
            "control_type": r.get("control_type") or control_type,
            "score": round(score_pct),
            "grade_5": score_to_grade5(r.get("total_score"), r.get("max_score") or 100),
            "date": _fmt_date(r.get("updated_at") or r.get("created_at"), tpl.get("date_format", "%d.%m.%Y")),
            "teacher": r.get("teacher_id") or "",
            "status": labels.get(status, status),
            "grade_id": r.get("grade_id"),
        })
    return out


def write_hemis_excel(rows: List[Dict[str, Any]], course: Optional[str] = None, group: Optional[str] = None, control_type: str = "JN") -> Path:
    tpl = load_hemis_template()
    data = hemis_rows(rows, control_type, tpl)
    wb = Workbook()
    ws = wb.active
    ws.title = tpl.get("sheet_name", "Baholar")[:31]
    columns = tpl["columns"]
    ws.append([c["header"] for c in columns])
    for d in data:
        ws.append([d.get(c["field"], "") for c in columns])
    _style_header(ws, len(columns))
    score_idx = next((i for i, c in enumerate(columns, start=1) if c["field"] == "score"), None)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = BORDER
        if score_idx:
            val = row[score_idx - 1].value
            if isinstance(val, (int, float)):
                row[score_idx - 1].fill = FAIL_FILL if val < config.settings.pass_threshold else GOOD_FILL if val >= config.settings.grade5_min else PatternFill()
    _autosize(ws)
    path = make_path("hemis", f"{course or 'all'}_{group or 'all'}", "xlsx")
    wb.save(path)
    return path


def write_hemis_csv(rows: List[Dict[str, Any]], course: Optional[str] = None, group: Optional[str] = None, control_type: str = "JN") -> Path:
    tpl = load_hemis_template()
    data = hemis_rows(rows, control_type, tpl)
    path = make_path("hemis", f"{course or 'all'}_{group or 'all'}", "csv")
    with open(path, "w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh, delimiter=";")
        writer.writerow([c["header"] for c in tpl["columns"]])
        for d in data:
            writer.writerow([d.get(c["field"], "") for c in tpl["columns"]])
    return path


# ---------------------------------------------------------------------------
# To'liq Excel hisobot
# ---------------------------------------------------------------------------
def write_grades_excel(rows: List[Dict[str, Any]], analytics: Optional[Dict[str, Any]] = None, course_name: Optional[str] = None, group_id: Optional[str] = None) -> Path:
    """Ko'p varaqli Excel: Baholar (mezonlar bilan), Statistika, Mavzular, Mezonlar, Talabalar."""
    wb = Workbook()

    # --- 1. Baholar ---
    ws = wb.active
    ws.title = "Baholar"
    criteria: List[str] = []
    for r in rows:
        for item in r.get("rubric") or []:
            if isinstance(item, dict) and item.get("name") and item["name"] not in criteria:
                criteria.append(item["name"])
    header = ["№", "Talaba ID", "F.I.Sh.", "Guruh", "Fan", "Mavzu", "Ball (100)", "Baho (5)", "AI balli", "Holat", "O'qituvchi",
              "Plagiat", "AI-matn", "O'xshashlik", "Sana"] + [f"{c} (ball)" for c in criteria] + ["Eng zaif mezon", "Fikr-mulohaza"]
    ws.append(header)
    for i, r in enumerate(rows, start=1):
        rubric = [it for it in (r.get("rubric") or []) if isinstance(it, dict)]
        by_name = {it.get("name"): it for it in rubric}
        weakest = None
        for it in rubric:
            if it.get("score") is None:
                continue
            ratio = percent(it.get("score"), it.get("max_score") or 1)
            if weakest is None or ratio < weakest[1]:
                weakest = (it.get("name"), ratio)
        pct = percent(r.get("total_score"), r.get("max_score") or 100)
        ws.append([
            i, r.get("student_id"), r.get("student_name") or "", r.get("group_name") or (group_id or ""), r.get("course") or (course_name or ""),
            r.get("topic") or "", round(pct, 1), score_to_grade5(r.get("total_score"), r.get("max_score") or 100),
            None if r.get("ai_total_score") is None else round(percent(r.get("ai_total_score"), r.get("max_score") or 100), 1),
            r.get("status") or ("confirmed" if r.get("confirmed") else "pending"), r.get("teacher_id") or "",
            r.get("plagiarism_score"), r.get("ai_likelihood"), r.get("similarity_score"),
            _fmt_date(r.get("updated_at") or r.get("created_at"), "%d.%m.%Y %H:%M"),
        ] + [(by_name.get(c) or {}).get("score") for c in criteria] + [weakest[0] if weakest else "", (r.get("feedback_summary") or "")[:500]])
    _style_header(ws, len(header))
    for row in ws.iter_rows(min_row=2):
        val = row[6].value
        if isinstance(val, (int, float)):
            row[6].fill = FAIL_FILL if val < config.settings.pass_threshold else GOOD_FILL if val >= config.settings.grade5_min else PatternFill()
    _autosize(ws)

    a = analytics or {}
    # --- 2. Statistika ---
    ws2 = wb.create_sheet("Statistika")
    ws2.append(["Ko'rsatkich", "Qiymat"])
    stats = [
        ("Fan", course_name or "Barcha"), ("Guruh", group_id or "Barcha"),
        ("Jami ishlar", a.get("total_grades", len(rows))), ("Talabalar soni", a.get("total_students", "")),
        ("O'rtacha ball (%)", a.get("average_score", "")), ("Mediana (%)", a.get("median_score", "")),
        ("O'zlashtirish (%)", a.get("pass_rate", "")), ("Sifat ko'rsatkichi (%)", a.get("quality_rate", "")),
        ("«5» soni", (a.get("distribution") or {}).get("5", "")), ("«4» soni", (a.get("distribution") or {}).get("4", "")),
        ("«3» soni", (a.get("distribution") or {}).get("3", "")), ("«2» soni", (a.get("distribution") or {}).get("2", "")),
        ("Tasdiqlangan", a.get("confirmed_count", "")), ("Kutilmoqda", a.get("pending_count", "")),
        ("Plagiat belgisi", (a.get("integrity") or {}).get("plagiarism_flags", "")), ("AI-matn belgisi", (a.get("integrity") or {}).get("ai_flags", "")),
        ("Tejalgan vaqt (soat)", a.get("time_saved_hours", "")), ("Hisobot sanasi", datetime.utcnow().strftime("%d.%m.%Y %H:%M")),
    ]
    for k, v in stats:
        ws2.append([k, v])
    _style_header(ws2, 2)
    _autosize(ws2)

    # --- 3. Mavzular ---
    ws3 = wb.create_sheet("Mavzular")
    ws3.append(["Mavzu", "O'rtacha ball (%)", "O'zlashtirish (%)", "Ishlar soni", "Qiyinchilik (%)"])
    for w in a.get("weak_topics") or []:
        ws3.append([w["topic"], w["average_score"], w["pass_rate"], w["count"], w["difficulty"]])
    _style_header(ws3, 5)
    _autosize(ws3)

    # --- 4. Mezonlar ---
    ws4 = wb.create_sheet("Mezonlar")
    ws4.append(["Mezon", "O'rtacha (%)", "Qiyinchilik (%)", "Baholar soni", "Chegaradan past (%)"])
    for c in a.get("criteria_stats") or []:
        ws4.append([c["criterion"], c["average_percent"], c["difficulty"], c["count"], c["below_pass_pct"]])
    _style_header(ws4, 5)
    _autosize(ws4)

    # --- 5. Talabalar ---
    ws5 = wb.create_sheet("Talabalar")
    ws5.append(["Talaba ID", "F.I.Sh.", "Guruh", "Ishlar", "O'rtacha (%)", "Oxirgi (%)", "Dinamika", "Baho (5)", "Xavf guruhi"])
    for st in a.get("student_stats") or []:
        ws5.append([st["student_id"], st.get("student_name") or "", st.get("group") or "", st["count"], st["average_score"], st["last_score"], st["trend"], st["grade_5"], "Ha" if st["at_risk"] else ""])
    _style_header(ws5, 9)
    _autosize(ws5)

    path = make_path("grades", f"{course_name or 'all'}_{group_id or 'all'}", "xlsx")
    wb.save(path)
    return path


# ---------------------------------------------------------------------------
# DOCX hujjatlar
# ---------------------------------------------------------------------------
def _doc_base(title: str, subtitle: Optional[str] = None) -> Document:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    h = doc.add_heading(title, level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if subtitle:
        p = doc.add_paragraph(subtitle)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return doc


def _table(doc: Document, header: List[str], rows: List[List[Any]]) -> None:
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    for i, h in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = str(h)
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
    for r in rows:
        cells = table.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = "" if v is None else str(v)


def render_syllabus_text(s: Dict[str, Any]) -> str:
    if not s:
        return ""
    lines = [f"{s.get('course')} — {s.get('weeks')} haftalik o'quv rejasi (jami {s.get('hours')} soat, {s.get('level')})", ""]
    hb = s.get("hours_breakdown") or {}
    lines.append(f"Ma'ruza: {hb.get('lecture')} s. | Amaliy: {hb.get('practice')} s. | Mustaqil ta'lim: {hb.get('independent')} s.")
    lines.append("")
    lines.append("Fanning maqsadi:")
    lines += [f"  • {g}" for g in s.get("goals", [])]
    lines.append("")
    lines.append("Kutilayotgan natijalar:")
    lines += [f"  • {g}" for g in s.get("learning_outcomes", [])]
    lines.append("")
    lines.append("Haftalik reja:")
    for w in s.get("weekly", []):
        lines.append(f"  {w['week']:>2}-hafta: {w['topic']}  [M:{w['lecture_hours']} A:{w['practice_hours']} MT:{w['independent_hours']}] — {w['assessment']}")
    lines.append("")
    lines.append("Baholash tizimi:")
    lines += [f"  • {p['type']} — {p['max_score']} ball: {p['description']}" for p in s.get("assessment_policy", [])]
    lines.append(f"  Shkala: {s.get('grade_scale', '')}")
    lines.append("")
    lines.append("Adabiyotlar:")
    lines += [f"  {i}. {lit}" for i, lit in enumerate(s.get("literature", []), start=1)]
    return "\n".join(lines)


def render_lesson_plan_text(lp: Dict[str, Any]) -> str:
    if not lp:
        return ""
    lines = [f"DARS REJASI — {lp.get('course')}", f"Mavzu: {lp.get('topic')} ({lp.get('week')}-hafta, {lp.get('duration_minutes')} daqiqa)", ""]
    lines.append("Dars maqsadlari:")
    lines += [f"  • {o}" for o in lp.get("objectives", [])]
    lines.append("")
    lines.append("Kompetensiyalar: " + ", ".join(lp.get("competencies", [])))
    lines.append("Metodlar: " + ", ".join(lp.get("methods", [])))
    lines.append("Jihozlar: " + ", ".join(lp.get("materials", [])))
    lines.append("")
    lines.append("Dars bosqichlari:")
    for st in lp.get("stages", []):
        lines.append(f"  {st['minutes']:>2} daq. — {st['stage']}: {st['activity']}")
    lines.append("")
    lines.append(f"Uyga vazifa: {lp.get('homework')}")
    lines.append(f"Baholash: {lp.get('assessment')}")
    return "\n".join(lines)


def render_tickets_text(tickets: List[Dict[str, Any]], course: str) -> str:
    if not tickets:
        return ""
    lines = []
    for t in tickets:
        lines.append(f"IMTIHON BILETI № {t['number']} — {course}")
        for i, q in enumerate(t["questions"], start=1):
            lines.append(f"  {i}. ({q['type']}) {q['text']}")
        lines.append("")
    return "\n".join(lines).rstrip()


def render_tests_text(questions: List[Dict[str, Any]], with_key: bool = True) -> str:
    if not questions:
        return ""
    lines = []
    for q in questions:
        lines.append(f"{q['number']}. {q['question']}")
        for k, v in q["options"].items():
            lines.append(f"   {k}) {v}")
        lines.append("")
    if with_key:
        lines.append("Javoblar kaliti: " + ", ".join(f"{q['number']}-{q['answer']}" for q in questions))
    return "\n".join(lines).rstrip()


def write_documents_docx(course_name: str, docs: Dict[str, Any]) -> Path:
    doc = _doc_base(f"{course_name}", "O'quv-uslubiy hujjatlar to'plami (avtomatik shakllantirilgan)")
    s = docs.get("syllabus")
    if s:
        doc.add_heading("1. Fan dasturi (silabus)", level=2)
        hb = s.get("hours_breakdown") or {}
        doc.add_paragraph(f"Daraja: {s.get('level')}. Jami: {s.get('hours')} soat ({s.get('weeks')} hafta). Ma'ruza — {hb.get('lecture')} s., amaliy — {hb.get('practice')} s., mustaqil ta'lim — {hb.get('independent')} s.")
        doc.add_paragraph("Fanning maqsadi:", style="Heading 3")
        for g in s.get("goals", []):
            doc.add_paragraph(g, style="List Bullet")
        doc.add_paragraph("Kutilayotgan o'quv natijalari:", style="Heading 3")
        for g in s.get("learning_outcomes", []):
            doc.add_paragraph(g, style="List Bullet")
        doc.add_paragraph("Haftalik taqvim-mavzu rejasi:", style="Heading 3")
        _table(doc, ["Hafta", "Mavzu", "Ma'ruza", "Amaliy", "Mustaqil", "Nazorat"],
               [[w["week"], w["topic"], w["lecture_hours"], w["practice_hours"], w["independent_hours"], w["assessment"]] for w in s.get("weekly", [])])
        doc.add_paragraph("Baholash tizimi:", style="Heading 3")
        _table(doc, ["Nazorat turi", "Ball", "Izoh"], [[p["type"], p["max_score"], p["description"]] for p in s.get("assessment_policy", [])])
        doc.add_paragraph(s.get("grade_scale", ""))
        doc.add_paragraph("Adabiyotlar:", style="Heading 3")
        for lit in s.get("literature", []):
            doc.add_paragraph(lit, style="List Number")
    lp = docs.get("lesson_plan")
    if lp:
        doc.add_page_break()
        doc.add_heading("2. Dars rejasi", level=2)
        doc.add_paragraph(f"Mavzu: {lp.get('topic')} ({lp.get('week')}-hafta, {lp.get('duration_minutes')} daqiqa)")
        doc.add_paragraph("Maqsadlar:", style="Heading 3")
        for o in lp.get("objectives", []):
            doc.add_paragraph(o, style="List Bullet")
        doc.add_paragraph("Metodlar: " + ", ".join(lp.get("methods", [])))
        doc.add_paragraph("Jihozlar: " + ", ".join(lp.get("materials", [])))
        _table(doc, ["Bosqich", "Daqiqa", "Faoliyat"], [[st["stage"], st["minutes"], st["activity"]] for st in lp.get("stages", [])])
        doc.add_paragraph(f"Uyga vazifa: {lp.get('homework')}")
        doc.add_paragraph(f"Baholash: {lp.get('assessment')}")
    tickets = docs.get("exam_tickets")
    if tickets:
        doc.add_page_break()
        doc.add_heading("3. Imtihon biletlari", level=2)
        for t in tickets:
            doc.add_paragraph(f"BILET № {t['number']}", style="Heading 3")
            for i, q in enumerate(t["questions"], start=1):
                doc.add_paragraph(f"{i}. ({q['type']}) {q['text']}")
            doc.add_paragraph("Kafedra mudiri: ____________     O'qituvchi: ____________")
    tests = docs.get("test_questions")
    if tests:
        doc.add_page_break()
        doc.add_heading("4. Test savollari", level=2)
        for q in tests:
            doc.add_paragraph(f"{q['number']}. {q['question']}")
            for k, v in q["options"].items():
                doc.add_paragraph(f"{k}) {v}", style="List Bullet")
        doc.add_paragraph("Javoblar kaliti: " + ", ".join(f"{q['number']}-{q['answer']}" for q in tests)).runs[0].bold = True
    path = make_path("hujjatlar", course_name, "docx")
    doc.save(path)
    return path


def write_syllabus_docx(course_name: str, syllabus: str, exam_ticket: str) -> str:
    """Eski API bilan moslik: matnli silabus va bilet."""
    doc = _doc_base(f"{course_name} — Silabus")
    for line in (syllabus or "").split("\n"):
        doc.add_paragraph(line)
    doc.add_page_break()
    doc.add_heading("Imtihon bileti", level=2)
    for line in (exam_ticket or "").split("\n"):
        doc.add_paragraph(line)
    path = make_path("syllabus", course_name, "docx")
    doc.save(path)
    return str(path)


def write_dean_report_docx(analytics: Dict[str, Any], narrative: str, course: Optional[str], group: Optional[str]) -> Path:
    doc = _doc_base("Tahliliy hisobot", f"Fan: {course or 'barcha'} | Guruh: {group or 'barcha'} | Sana: {datetime.utcnow().strftime('%d.%m.%Y')}")
    doc.add_heading("1. Umumiy ko'rsatkichlar", level=2)
    dist = analytics.get("distribution") or {}
    _table(doc, ["Ko'rsatkich", "Qiymat"], [
        ["Jami ishlar", analytics.get("total_grades", 0)], ["Talabalar", analytics.get("total_students", 0)],
        ["O'rtacha ball (%)", analytics.get("average_score", 0)], ["O'zlashtirish (%)", analytics.get("pass_rate", 0)],
        ["Sifat ko'rsatkichi (%)", analytics.get("quality_rate", 0)],
        ["Baholar taqsimoti", f"5: {dist.get('5', 0)} | 4: {dist.get('4', 0)} | 3: {dist.get('3', 0)} | 2: {dist.get('2', 0)}"],
        ["Tejalgan vaqt (soat)", analytics.get("time_saved_hours", 0)],
    ])
    doc.add_heading("2. Mavzular bo'yicha tahlil", level=2)
    _table(doc, ["Mavzu", "O'rtacha (%)", "O'zlashtirish (%)", "Ishlar"], [[w["topic"], w["average_score"], w["pass_rate"], w["count"]] for w in analytics.get("weak_topics") or []])
    doc.add_heading("3. Mezonlar bo'yicha tahlil", level=2)
    _table(doc, ["Mezon", "O'rtacha (%)", "Chegaradan past (%)"], [[c["criterion"], c["average_percent"], c["below_pass_pct"]] for c in analytics.get("criteria_stats") or []])
    doc.add_heading("4. Xavf guruhidagi talabalar", level=2)
    risk = analytics.get("at_risk_students") or []
    if risk:
        _table(doc, ["Talaba", "Guruh", "O'rtacha (%)", "Oxirgi (%)"], [[r.get("student_name") or r["student_id"], r.get("group") or "", r["average_score"], r["last_score"]] for r in risk])
    else:
        doc.add_paragraph("Xavf guruhidagi talabalar yo'q.")
    doc.add_heading("5. Tahliliy xulosa va tavsiyalar", level=2)
    for para in (narrative or "").split("\n\n"):
        if para.strip():
            doc.add_paragraph(para.strip())
    doc.add_paragraph("\nKafedra mudiri: ____________          O'qituvchi: ____________")
    path = make_path("hisobot", f"{course or 'all'}_{group or 'all'}", "docx")
    doc.save(path)
    return path


def write_student_feedback_docx(grade: Dict[str, Any]) -> Path:
    """Bitta talaba uchun fikr-mulohaza varaqasi."""
    doc = _doc_base("Baholash natijasi va fikr-mulohaza", f"{grade.get('course') or ''} — {grade.get('topic') or ''}")
    _table(doc, ["Maydon", "Qiymat"], [
        ["Talaba", f"{grade.get('student_name') or ''} ({grade.get('student_id') or ''})"], ["Guruh", grade.get("group_name") or ""],
        ["Ball", f"{round(grade.get('total_score') or 0)} / {round(grade.get('max_score') or 100)}"],
        ["Baho (5 ballik)", f"{grade.get('grade_5')} ({GRADE_LABELS.get(grade.get('grade_5') or 2, '')})"],
        ["Holat", grade.get("status") or ""], ["Sana", _fmt_date(grade.get("updated_at"), "%d.%m.%Y")],
    ])
    doc.add_heading("Mezonlar bo'yicha", level=2)
    _table(doc, ["Mezon", "Ball", "Maks.", "Asos"], [[r.get("name"), r.get("score"), r.get("max_score"), r.get("evidence") or ""] for r in grade.get("rubric") or []])
    fb = grade.get("feedback") or {}
    doc.add_heading("Umumiy xulosa", level=2)
    doc.add_paragraph(fb.get("summary") or "")
    for title, key in (("Kuchli tomonlar", "strengths"), ("Aniqlangan xatolar", "errors"), ("Tavsiyalar", "suggestions")):
        items = fb.get(key) or []
        if items:
            doc.add_heading(title, level=3)
            for it in items:
                doc.add_paragraph(str(it), style="List Bullet")
    if grade.get("teacher_comment"):
        doc.add_heading("O'qituvchi izohi", level=3)
        doc.add_paragraph(grade["teacher_comment"])
    path = make_path("feedback", f"{grade.get('student_id') or 'student'}", "docx")
    doc.save(path)
    return path


def rows_to_csv_bytes(rows: List[Dict[str, Any]]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";")
    writer.writerow(["student_id", "student_name", "group_name", "course", "topic", "score", "grade_5", "status", "date"])
    for r in rows:
        writer.writerow([r.get("student_id"), r.get("student_name") or "", r.get("group_name") or "", r.get("course") or "", r.get("topic") or "",
                         round(percent(r.get("total_score"), r.get("max_score") or 100)), score_to_grade5(r.get("total_score"), r.get("max_score") or 100),
                         r.get("status") or "", _fmt_date(r.get("updated_at"), "%d.%m.%Y")])
    return buf.getvalue().encode("utf-8-sig")
