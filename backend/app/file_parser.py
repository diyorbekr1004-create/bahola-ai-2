"""Yuklangan fayllardan matn ajratib olish: TXT/MD, DOCX, PDF, ZIP, XLSX/CSV.

Barcha funksiyalar baytlar bilan ishlaydi (diskka yozilmaydi).
"""
from __future__ import annotations

import csv
import io
import re
import zipfile
from pathlib import PurePosixPath
from typing import Iterable, Iterator, List, Optional, Tuple

TEXT_EXTENSIONS = {".txt", ".md", ".text", ".rtf"}
DOCX_EXTENSIONS = {".docx"}
PDF_EXTENSIONS = {".pdf"}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | DOCX_EXTENSIONS | PDF_EXTENSIONS


class UnsupportedFileError(ValueError):
    pass


def _ext(filename: Optional[str]) -> str:
    if not filename:
        return ""
    return PurePosixPath(filename.replace("\\", "/")).suffix.lower()


def _decode_bytes(raw: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "cp1251", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _docx_text(raw: bytes) -> str:
    from docx import Document  # lazy import

    doc = Document(io.BytesIO(raw))
    parts: List[str] = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _pdf_text(raw: bytes) -> str:
    from pypdf import PdfReader  # lazy import

    reader = PdfReader(io.BytesIO(raw))
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:  # pragma: no cover - buzilgan sahifa
            pages.append("")
    return "\n".join(pages)


def normalize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_text(filename: Optional[str], raw: bytes) -> str:
    """Fayl kengaytmasiga qarab matn qaytaradi. Noma'lum tur -> UTF-8 sifatida o'qishga urinadi."""
    ext = _ext(filename)
    if ext in DOCX_EXTENSIONS:
        return normalize_text(_docx_text(raw))
    if ext in PDF_EXTENSIONS:
        return normalize_text(_pdf_text(raw))
    if ext in TEXT_EXTENSIONS or ext == "":
        return normalize_text(_decode_bytes(raw))
    # Kengaytma noma'lum — magic bytes bo'yicha urinib ko'ramiz
    if raw[:4] == b"%PDF":
        return normalize_text(_pdf_text(raw))
    if raw[:2] == b"PK":
        try:
            return normalize_text(_docx_text(raw))
        except Exception:
            pass
    if b"\x00" in raw[:1024]:
        raise UnsupportedFileError(f"Qo'llab-quvvatlanmaydigan fayl turi: {ext or 'noma`lum'}")
    return normalize_text(_decode_bytes(raw))


def word_count(text: str) -> int:
    return len(re.findall(r"\w+", text or ""))


def is_zip(filename: Optional[str], raw: bytes) -> bool:
    return _ext(filename) == ".zip" or (raw[:2] == b"PK" and _ext(filename) not in DOCX_EXTENSIONS and zipfile.is_zipfile(io.BytesIO(raw)) and _ext(filename) != ".xlsx")


def iter_zip_files(raw: bytes) -> Iterator[Tuple[str, bytes]]:
    """ZIP ichidagi qo'llab-quvvatlanadigan fayllarni (nom, bayt) ko'rinishida qaytaradi."""
    with zipfile.ZipFile(io.BytesIO(raw)) as zf:
        for info in zf.infolist():
            if info.is_dir():
                continue
            name = info.filename
            base = PurePosixPath(name).name
            if base.startswith("._") or "__MACOSX" in name:
                continue
            if _ext(name) not in SUPPORTED_EXTENSIONS:
                continue
            yield base, zf.read(info)


def student_id_from_filename(filename: str) -> str:
    """`AT2101_Aliyev_Vali.docx` -> `AT2101_Aliyev_Vali` (kengaytmasiz, tozalangan)."""
    stem = PurePosixPath(filename.replace("\\", "/")).stem
    stem = re.sub(r"[^\w\-\.]+", "_", stem).strip("_")
    return stem or "unknown"


# ---------------------------------------------------------------------------
# Talabalar ro'yxatini import qilish (XLSX / CSV)
# ---------------------------------------------------------------------------
_ID_KEYS = ("student_id", "hemis_id", "hemis id", "talaba id", "talaba_id", "id", "№ talaba", "id raqami")
_NAME_KEYS = ("full_name", "fullname", "f.i.sh.", "f.i.sh", "fish", "ism", "ism familiya", "familiya ism", "name", "talaba", "talaba f.i.sh.")
_GROUP_KEYS = ("group_name", "group", "guruh", "guruh nomi", "gurux")


def _pick(row: dict, keys: Iterable[str]) -> Optional[str]:
    lowered = {str(k).strip().lower(): v for k, v in row.items() if k is not None}
    for k in keys:
        if k in lowered and lowered[k] is not None and str(lowered[k]).strip() != "" and str(lowered[k]).lower() != "nan":
            return str(lowered[k]).strip()
    return None


def parse_student_list(filename: Optional[str], raw: bytes) -> List[dict]:
    """XLSX yoki CSV dan [{student_id, full_name, group_name}] ro'yxatini qaytaradi.

    Ustun nomlari moslashuvchan: `student_id`/`HEMIS ID`/`Talaba ID`, `F.I.Sh.`/`full_name`, `Guruh`/`group`.
    """
    ext = _ext(filename)
    rows: List[dict] = []
    if ext in {".xlsx", ".xlsm", ".xls"}:
        import pandas as pd

        df = pd.read_excel(io.BytesIO(raw))
        rows = df.to_dict(orient="records")
    else:
        text = _decode_bytes(raw)
        sample = text[:2048]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(io.StringIO(text), dialect=dialect)
        rows = list(reader)

    out: List[dict] = []
    for r in rows:
        sid = _pick(r, _ID_KEYS)
        name = _pick(r, _NAME_KEYS)
        group = _pick(r, _GROUP_KEYS)
        if not sid and not name:
            continue
        if sid and sid.endswith(".0"):
            sid = sid[:-2]
        out.append({"student_id": sid, "full_name": name, "group_name": group})
    return out
