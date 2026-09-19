import os
from datetime import datetime
from docx import Document
import pandas as pd


EXPORT_DIR = os.path.join(os.getcwd(), "exports")
os.makedirs(EXPORT_DIR, exist_ok=True)


def write_syllabus_docx(course_name: str, syllabus: str, exam_ticket: str) -> str:
    name = f"syllabus_{course_name.replace(' ', '_')}_{datetime.utcnow().strftime('%Y%m%d%H%M%S')}.docx"
    path = os.path.join(EXPORT_DIR, name)
    doc = Document()
    doc.add_heading(f"{course_name} — Silabus", level=1)
    for line in syllabus.split('\n'):
        doc.add_paragraph(line)
    doc.add_page_break()
    doc.add_heading("Imtihon bileti", level=2)
    for line in exam_ticket.split('\n'):
        doc.add_paragraph(line)
    doc.save(path)
    return path


def write_grades_excel(rows: list, course_name: str = "grades", group_id: str = None) -> str:
    export_rows = []
    for row in rows:
        rubric = row.get("rubric") or []
        weakest = None
        if rubric:
            for item in rubric:
                if isinstance(item, dict):
                    score = item.get("score")
                    max_score = item.get("max_score") or 1
                    if score is None:
                        continue
                    ratio = float(score) / float(max_score)
                    candidate = {"name": item.get("name", "Unknown"), "score": score, "max_score": max_score, "ratio": ratio}
                    if weakest is None or candidate["ratio"] < weakest["ratio"]:
                        weakest = candidate
        export_rows.append({
            "student_id": row.get("student_id"),
            "course": row.get("course") or course_name,
            "group_id": group_id or "",
            "grade_id": row.get("grade_id"),
            "submission_id": row.get("submission_id"),
            "total_score": row.get("total_score"),
            "confirmed": row.get("confirmed"),
            "updated_at": row.get("updated_at"),
            "weakest_topic": weakest["name"] if weakest else "",
            "weakest_topic_score": weakest["score"] if weakest else "",
            "weakest_topic_max_score": weakest["max_score"] if weakest else "",
        })

    name = f"grades_{course_name.replace(' ', '_')}_{datetime.utcnow().strftime('%Y%m%d%H%M%S')}.xlsx"
    path = os.path.join(EXPORT_DIR, name)
    df = pd.DataFrame(export_rows)
    df.to_excel(path, index=False)
    return path
