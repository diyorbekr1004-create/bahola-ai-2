import io
import zipfile

import pytest
from docx import Document
from pypdf import PdfWriter

from app.file_parser import UnsupportedFileError, extract_text, is_zip, iter_zip_files, parse_student_list, student_id_from_filename, word_count
from app.plagiarism_detector import analyze_text, find_most_similar, text_similarity


def test_extract_txt_variants():
    assert extract_text("a.txt", "Salom dunyo".encode("utf-8")) == "Salom dunyo"
    assert "Привет" in extract_text("a.txt", "Привет".encode("cp1251"))
    assert extract_text("a.md", b"\xef\xbb\xbfBOM matn") == "BOM matn"


def test_extract_docx_and_pdf():
    doc = Document()
    doc.add_paragraph("Birinchi abzats.")
    t = doc.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "A", "B"
    buf = io.BytesIO()
    doc.save(buf)
    out = extract_text("ish.docx", buf.getvalue())
    assert "Birinchi abzats." in out and "A | B" in out

    w = PdfWriter()
    w.add_blank_page(width=200, height=200)
    pbuf = io.BytesIO()
    w.write(pbuf)
    assert extract_text("bo'sh.pdf", pbuf.getvalue()) == ""


def test_unsupported_binary():
    with pytest.raises(UnsupportedFileError):
        extract_text("x.bin", b"\x00\x01\x02\x03" * 10)


def test_zip_iteration_and_filename_ids():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("guruh/AT-21 Aliyev.txt", "matn")
        zf.writestr("__MACOSX/._x.txt", "junk")
        zf.writestr("img.png", b"png")
    raw = buf.getvalue()
    assert is_zip("ishlar.zip", raw)
    files = list(iter_zip_files(raw))
    assert files == [("AT-21 Aliyev.txt", b"matn")]
    assert student_id_from_filename("AT-21 Aliyev.txt") == "AT-21_Aliyev"
    assert word_count("bir ikki uch") == 3


def test_parse_student_list_csv_flexible_headers():
    rows = parse_student_list("t.csv", "Talaba ID,Ism familiya,Guruh\n1001,Ali Valiyev,G1\n".encode())
    assert rows == [{"student_id": "1001", "full_name": "Ali Valiyev", "group_name": "G1"}]


def test_similarity_and_analysis():
    a = "Bu matn birinchi talaba tomonidan yozilgan va u texnologiya haqida."
    assert text_similarity(a, a) >= 0.99
    assert text_similarity(a, "Butunlay boshqa mavzu: kimyoviy reaksiyalar tezligi haqida.") < 0.3
    best, label = find_most_similar(a, [("x", "hech qanday aloqasi yo'q"), ("y", a + " qo'shimcha")])
    assert label == "y" and best > 0.6
    out = analyze_text("Bu oddiy matn. Bu oddiy matn. Bu oddiy matn.")
    assert out["plagiarism_score"] > 0.3 and "plagiarism" in out["flags"]
    out2 = analyze_text("Bu matn har bir jumlasi noyob va turlicha yozilgan namunadir. Har bir jumla boshqa mavzuni ochib beradi.")
    assert out2["plagiarism_score"] <= 0.5
