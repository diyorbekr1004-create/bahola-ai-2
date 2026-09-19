import io

from openpyxl import load_workbook


def _seed(client, good_text, weak_text):
    for i, (sid, text) in enumerate([("R-1", good_text), ("R-2", weak_text), ("R-3", good_text + " Yana bir fikr."), ("R-4", weak_text + " Kam.")]):
        r = client.post("/api/grade", data={"text": text, "student_id": sid, "student_name": f"Talaba {i}", "group_name": "RP-01", "course": "Hisobot fani", "topic": "Mavzu A" if i % 2 == 0 else "Mavzu B"})
        assert r.status_code == 200


def test_reports_analytics_and_exports(client, good_text, weak_text, auth_headers):
    _seed(client, good_text, weak_text)
    r = client.post("/api/reports", json={"course_name": "Hisobot fani", "group_name": "RP-01", "export_format": "all", "control_type": "ON"})
    assert r.status_code == 200, r.text
    j = r.json()
    a = j["analytics"]
    assert a["total_grades"] == 4 and a["total_students"] == 4
    assert 0 <= a["pass_rate"] <= 100 and 0 <= a["quality_rate"] <= 100
    assert sum(a["distribution"].values()) == 4
    assert len(a["weak_topics"]) == 2 and a["weak_topics"][0]["difficulty"] >= a["weak_topics"][1]["difficulty"]
    assert a["criteria_stats"] and a["group_stats"][0]["group"] == "RP-01"
    assert a["time_saved_hours"] > 0
    assert "Hisobot fani" in j["summary"]
    assert "o'zlashtirish" in j["narrative"].lower()

    for key in ("download_url", "hemis_download_url", "csv_download_url", "dean_report_url"):
        assert j[key], key
        dl = client.get(j[key])
        assert dl.status_code == 200, key

    # HEMIS jadval ustunlari shablonga mos
    wb = load_workbook(io.BytesIO(client.get(j["hemis_download_url"]).content))
    ws = wb["Baholar"]
    headers = [c.value for c in ws[1]]
    assert headers[:3] == ["№", "Talaba ID (HEMIS)", "F.I.Sh."]
    assert "Baho (5)" in headers and "Nazorat turi" in headers
    assert ws.cell(row=2, column=headers.index("Nazorat turi") + 1).value == "ON"
    assert ws.cell(row=2, column=headers.index("Baho (5)") + 1).value in (2, 3, 4, 5)

    # To'liq Excel — ko'p varaqli
    wb2 = load_workbook(io.BytesIO(client.get(j["download_url"]).content))
    assert {"Baholar", "Statistika", "Mavzular", "Mezonlar", "Talabalar"} <= set(wb2.sheetnames)


def test_reports_export_endpoint_direct(client, auth_headers):
    r = client.get("/api/reports/export", params={"format": "hemis", "course": "Hisobot fani"}, headers=auth_headers)
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    r2 = client.get("/api/reports/export", params={"format": "csv", "course": "Hisobot fani"})
    assert r2.status_code == 200 and "Talaba ID" in r2.content.decode("utf-8-sig")
    r3 = client.get("/api/reports/export", params={"format": "dean", "course": "Hisobot fani"})
    assert r3.status_code == 200 and r3.content[:2] == b"PK"
    r4 = client.get("/api/reports/export", params={"format": "excel", "course": "Mavjud emas"})
    assert r4.status_code == 404


def test_reports_empty_scope(client):
    r = client.post("/api/reports", json={"course_name": "Yo'q fan"})
    assert r.status_code == 200
    assert r.json()["analytics"]["total_grades"] == 0
    assert r.json()["download_url"] is None
