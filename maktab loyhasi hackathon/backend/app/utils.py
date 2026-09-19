import os
from datetime import datetime
from docx import Document
import pandas as pd


EXPORT_DIR = os.path.join(os.getcwd(), "exports")
os.makedirs(EXPORT_DIR, exist_ok=True)


def write_syllabus_docx(course_name: str, syllabus: str, exam_ticket: str) -> str:
    name = f"syllabus_{course_name.replace(' ', '_')}_{datetime.utcnow().strftime('%Y%m%d%H%M%S')}.docx"
    path = os.path.join(EXPORT_DIR, name)
    doc = Document()
    doc.add_heading(f"{course_name} — Silabus", level=1)
    for line in syllabus.split('\n'):
        doc.add_paragraph(line)
    doc.add_page_break()
    doc.add_heading("Imtihon bileti", level=2)
    for line in exam_ticket.split('\n'):
        doc.add_paragraph(line)
    doc.save(path)
    return path


def write_grades_excel(rows: list, course_name: str = "grades") -> str:
    name = f"grades_{course_name.replace(' ', '_')}_{datetime.utcnow().strftime('%Y%m%d%H%M%S')}.xlsx"
    path = os.path.join(EXPORT_DIR, name)
    df = pd.DataFrame(rows)
    df.to_excel(path, index=False)
    return path
